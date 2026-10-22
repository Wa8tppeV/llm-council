"""Offline upload and council request smoke tests; no model API charges."""
import io
import tempfile
import unittest
from unittest.mock import AsyncMock, patch

import pymupdf
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from PIL import Image

from backend import attachments, storage
from backend.main import app, prepare, SendMessageRequest
from backend.council import choose_models


class FileUploadTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.data1 = patch.object(attachments, "DATA_DIR", self.temp.name)
        self.data2 = patch.object(storage, "DATA_DIR", self.temp.name)
        self.data1.start()
        self.data2.start()
        self.addCleanup(self.data1.stop)
        self.addCleanup(self.data2.stop)
        self.client = TestClient(app)
        self.conversation = self.client.post("/api/conversations", json={}).json()["id"]
        self.route = f"/api/conversations/{self.conversation}/files"

    def upload(self, name, content):
        return self.client.post(self.route, files={"file": (name, content)})

    def test_text_and_csv(self):
        x = self.upload("notlar.txt", "beton sinifi C30".encode())
        self.assertEqual(x.status_code, 200)
        fid = x.json()["id"]
        text, files = attachments.context_for_question(
            self.conversation, [fid], "beton sinifi ne?"
        )
        self.assertIn("C30", text)
        self.assertEqual(files[0]["name"], "notlar.txt")
        self.assertEqual(self.upload("metraj.csv", b"kalip,100\nbeton,25").status_code, 200)

    def test_pdf_docx_xlsx(self):
        pdf = pymupdf.open()
        page = pdf.new_page()
        page.insert_text((50, 70), "PDF C30 BETON")
        self.assertEqual(self.upload("rapor.pdf", pdf.tobytes()).status_code, 200)
        pdf.close()
        doc = Document()
        doc.add_paragraph("Word sozlesme - isveren")
        buf = io.BytesIO()
        doc.save(buf)
        self.assertEqual(self.upload("sozlesme.docx", buf.getvalue()).status_code, 200)
        excel = Workbook()
        excel.active["A1"] = "Demir"
        excel.active["B1"] = 120.5
        buf = io.BytesIO()
        excel.save(buf)
        response = self.upload("metraj.xlsx", buf.getvalue())
        self.assertEqual(response.status_code, 200, response.text)
        content, _ = attachments.context_for_question(
            self.conversation, [response.json()["id"]], "Demir kac ton?"
        )
        self.assertIn("120.5", content)

    def test_image_vision_and_private_storage(self):
        img = Image.new("RGB", (32, 32), color="white")
        buf = io.BytesIO()
        img.save(buf, format="PNG")
        with patch("backend.attachments._vision", new_callable=AsyncMock) as vision:
            vision.return_value = "Dikdortgen gorunuyor"
            response = self.upload("plan.png", buf.getvalue())
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(vision.await_count, 1)
        data = (attachments._folder(self.conversation) /
                (response.json()["id"] + ".json")).read_text(encoding="utf-8")
        self.assertIn("Dikdortgen", data)
        self.assertNotIn("iVBOR", data)

    def test_reject_types_and_limits(self):
        self.assertEqual(self.upload("plan.dwg", b"DWG data").status_code, 415)
        self.assertEqual(self.upload("huge.txt", b"a" * (12 * 1024 * 1024 + 1)).status_code, 413)
        self.assertEqual(self.upload("bad.pdf", b"not a pdf").status_code, 422)

    def test_isolation_and_source_selection(self):
        x = self.upload("project.txt", b"Foundation: 25 m3")
        fid = x.json()["id"]
        other = self.client.post("/api/conversations", json={}).json()["id"]
        response = self.client.post(
            f"/api/conversations/{other}/message",
            json={"content": "Foundation?", "file_ids": [fid]}
        )
        self.assertEqual(response.status_code, 404)
        q, labels = prepare(self.conversation, SendMessageRequest(content="Foundation?", file_ids=[fid]))
        self.assertIn("25 m3", q)
        self.assertIn("GUVENILMEZ VERIDIR", q)
        self.assertEqual(len(labels), 1)


class CouncilModeTests(unittest.IsolatedAsyncioTestCase):
    async def test_deep_uses_every_installed_local_model(self):
        status = {
            "connected": True,
            "selected_models": ["qwen3:8b", "gemma3:4b"],
            "installed_models": ["qwen3:8b", "gemma3:4b", "qwen3-coder:30b"]
        }
        with patch("backend.council.ollama_status", new_callable=AsyncMock) as fake:
            fake.return_value = status
            self.assertEqual(len(await choose_models(False)), 2)
            self.assertEqual(len(await choose_models(True)), 3)


if __name__ == "__main__":
    unittest.main()
