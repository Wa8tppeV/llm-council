"""Private local attachment extraction and query-relevant excerpts.

Only derived text is retained under data/attachments; uploaded binary bytes are
processed in memory. Ollama image analysis happens on localhost.
"""
import base64
import io
import json
import zipfile
import re
import uuid
from pathlib import Path

import pymupdf
from docx import Document
from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook
from PIL import Image, ImageOps, UnidentifiedImageError

from .config import DATA_DIR
from .ollama import ollama_status, query_model

MAX_FILE_BYTES = 12 * 1024 * 1024
MAX_TEXT_CHARS = 140000
ALLOWED = {".pdf", ".docx", ".xlsx", ".csv", ".txt", ".md",
           ".png", ".jpg", ".jpeg", ".webp"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp"}


def _folder(conversation_id: str) -> Path:
    # Never turn a client supplied ID into an arbitrary path.
    try:
        checked = str(uuid.UUID(conversation_id))
    except (ValueError, AttributeError) as exc:
        raise HTTPException(status_code=400, detail="Invalid conversation id") from exc
    path = Path(DATA_DIR) / "attachments" / checked
    path.mkdir(parents=True, exist_ok=True)
    return path


def _short(text: str, limit: int = MAX_TEXT_CHARS) -> str:
    if len(text) <= limit:
        return text
    return text[:limit] + "\n[UYARI: Dosyanin kalan metni bu surumde sinirlandi.]"


async def _vision(image: Image.Image, instruction: str) -> str:
    status = await ollama_status()
    visual = next(
        (m for m in status.get("installed_models", [])
         if m.startswith(("gemma3:4b", "qwen2.5vl:", "qwen3-vl:"))),
        None,
    )
    if not visual:
        return "[Gorsel yorumlanamadi: Ollama'da gemma3:4b gibi bir goruntu modeli bulunamadi.]"
    image = ImageOps.exif_transpose(image).convert("RGB")
    image.thumbnail((1500, 1500))
    stream = io.BytesIO()
    image.save(stream, format="JPEG", quality=82)
    blob = base64.b64encode(stream.getvalue()).decode("ascii")
    response = await query_model(
        visual,
        [{"role": "user", "content": instruction, "images": [blob]}],
        timeout=240.0,
    )
    return (response or {}).get("content") or "[Gorsel analizi basarisiz oldu.]"


async def _extract_pdf(raw: bytes) -> str:
    segments = []
    with pymupdf.open(stream=raw, filetype="pdf") as pdf:
        total = len(pdf)
        for number, page in enumerate(pdf, 1):
            if number > 80:
                segments.append("[UYARI: PDF 80 sayfadan sonra islenmedi.]")
                break
            content = page.get_text("text").strip()
            if not content and number <= 2:
                pix = page.get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5), alpha=False)
                with Image.open(io.BytesIO(pix.tobytes("png"))) as img:
                    content = await _vision(img,
                        "Read and accurately describe this scanned document page in Turkish. "
                        "Transcribe visible important figures and titles. Do not invent missing numbers.")
            elif not content:
                content = "[Bu sayfa taranmis, metin okunamadi.]"
            segments.append(f"[PDF sayfa {number}/{total}]\n{content}")
    return _short("\n\n".join(segments))


def _extract_word(raw: bytes) -> str:
    doc = Document(io.BytesIO(raw))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    for index, table in enumerate(doc.tables, 1):
        paragraphs.append(f"[Tablo {index}]")
        for row in list(table.rows)[:1500]:
            paragraphs.append(" | ".join(cell.text.replace("\n", " ") for cell in row.cells))
    return _short("\n".join(paragraphs))


def _extract_excel(raw: bytes) -> str:
    book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    parts = []
    used = 0
    try:
        for sheet in book.worksheets[:25]:
            parts.append(f"[EXCEL SAYFA: {sheet.title}; satir={sheet.max_row}, sutun={sheet.max_column}]")
            for idx, row in enumerate(sheet.iter_rows(values_only=True), 1):
                if idx > 1200:
                    parts.append("[UYARI: Bu sayfanin 1200 satirdan sonrasi alinmadi.]")
                    break
                cells = [f"{j+1}:{str(value)[:220]}" for j, value in enumerate(row[:32]) if value is not None]
                if cells:
                    line = f"Satir {idx} | " + " | ".join(cells)
                    parts.append(line)
                    used += len(line) + 1
                if used > MAX_TEXT_CHARS:
                    return _short("\n".join(parts))
    finally:
        book.close()
    return _short("\n".join(parts))


def _validate_office_zip(raw: bytes) -> None:
    # Avoid ZIP-bomb DOCX/XLSX archives: limit inflated bytes and archive entries.
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        entries = archive.infolist()
        if len(entries) > 2000 or sum(e.file_size for e in entries) > 60 * 1024 * 1024:
            raise ValueError("Office archive exceeds safe extraction limits")


async def save_upload(conversation_id: str, uploaded: UploadFile) -> dict:
    name = Path((uploaded.filename or "isimsiz").replace("\\", "/")).name[:110]
    suffix = Path(name).suffix.lower()
    if suffix not in ALLOWED:
        raise HTTPException(
            status_code=415,
            detail="Desteklenmeyen dosya. PDF, DOCX, XLSX, CSV, TXT, JPG, PNG veya WEBP kullanin.",
        )
    raw = await uploaded.read(MAX_FILE_BYTES + 1)
    await uploaded.close()
    if not raw:
        raise HTTPException(status_code=400, detail="Dosya bos.")
    if len(raw) > MAX_FILE_BYTES:
        raise HTTPException(status_code=413, detail="Dosya en fazla 12 MB olabilir.")
    try:
        if suffix == ".pdf":
            if not raw.startswith(b"%PDF"):
                raise ValueError("PDF signature invalid")
            text = await _extract_pdf(raw)
        elif suffix == ".docx":
            _validate_office_zip(raw)
            text = _extract_word(raw)
        elif suffix == ".xlsx":
            _validate_office_zip(raw)
            text = _extract_excel(raw)
        elif suffix in IMAGE_EXT:
            with Image.open(io.BytesIO(raw)) as img:
                img.verify()
            with Image.open(io.BytesIO(raw)) as img:
                text = "[Gorsel kaynak: " + name + "]\n" + await _vision(
                    img,
                    "Bu fotografin/teknik cizimin gorunen detaylarini Turkce incele. "
                    "Basliklar, olculer, sayilar ve olasi belirsizlikleri ayri belirt. "
                    "Okunamayan rakamlari uydurma."
                )
        else:
            text = raw.decode("utf-8-sig", errors="replace")
    except (Exception,) as exc:
        # Never return private file bytes, stack traces or paths to the browser.
        raise HTTPException(status_code=422, detail=f"Dosya acilamadi veya okunamadi: {type(exc).__name__}") from exc

    item = {
        "id": str(uuid.uuid4()),
        "name": name,
        "kind": suffix,
        "bytes": len(raw),
        "text": _short(text),
    }
    path = _folder(conversation_id) / (item["id"] + ".json")
    path.write_text(json.dumps(item, ensure_ascii=False), encoding="utf-8")
    return {"id": item["id"], "name": name, "kind": suffix,
            "bytes": len(raw), "characters": len(item["text"]),
            "warning": "Gorsel dogrulama gerektirebilir" if suffix in IMAGE_EXT else None}


def _load(conversation_id: str, file_id: str) -> dict:
    try:
        valid_id = str(uuid.UUID(file_id))
    except (ValueError, AttributeError) as exc:
        raise HTTPException(status_code=400, detail="Invalid file id") from exc
    path = _folder(conversation_id) / (valid_id + ".json")
    if not path.is_file():
        raise HTTPException(status_code=404, detail="Ek dosya bulunamadi.")
    return json.loads(path.read_text(encoding="utf-8"))


def _relevant_excerpt(text: str, question: str, budget: int) -> tuple[str, bool]:
    if len(text) <= budget:
        return text, False
    # Split by source paragraphs into bounded chunks, score by question words.
    pieces = []
    for p in text.splitlines():
        for offset in range(0, len(p), 850):
            snippet = p[offset:offset + 850]
            if snippet.strip():
                pieces.append(snippet)
    terms = set(re.findall(r"[\wÇĞİÖŞÜçğıöşü]{4,}", question.casefold()))
    scores = []
    for i, p in enumerate(pieces):
        lower = p.casefold()
        score = sum(min(lower.count(t), 3) for t in terms)
        scores.append((score, i, p))
    top = sorted(scores, key=lambda x: (-x[0], x[1]))
    indexes = {0}
    count = len(pieces[0]) if pieces else 0
    for _, idx, p in top:
        if idx in indexes:
            continue
        if count + len(p) + 15 > budget:
            continue
        indexes.add(idx)
        count += len(p) + 15
    result = "\n".join(f"[Bolum {i+1}] {pieces[i]}" for i in sorted(indexes))
    return result, True


def context_for_question(conversation_id: str, file_ids: list[str], question: str) -> tuple[str, list[dict]]:
    if len(file_ids) > 6:
        raise HTTPException(status_code=400, detail="Tek soruda en fazla 6 belge.")
    if len(set(file_ids)) != len(file_ids):
        raise HTTPException(status_code=400, detail="Tekrarlanan dosya secimi.")
    items = [_load(conversation_id, file_id) for file_id in file_ids]
    excerpts = []
    budget = 15000
    for item in items:
        available = max(350, min(8500, budget // max(1, len(items) - len(excerpts))))
        part, shortened = _relevant_excerpt(item["text"], question, available)
        excerpts.append(f"[DOSYA: {item['name']}]\n{part}\n" +
                        ("[UYARI: Yalnizca soru ile ilgili secilmis bolumler gosteriliyor. "
                         "Tum dosya taranmadi.]" if shortened else ""))
        budget -= min(available, len(part))
    files = [{"id": x["id"], "name": x["name"], "kind": x["kind"]} for x in items]
    return "\n\n".join(excerpts), files
